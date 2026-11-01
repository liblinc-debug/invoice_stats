#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from pypdf import PdfReader


@dataclass(frozen=True)
class Invoice:
    category: str
    source: Path
    invoice_date: date
    number: str
    item_name: str
    amount: Decimal


def extract_pdf_text(pdf_path: Path) -> str:
    reader = PdfReader(str(pdf_path))
    text = normalize_pdf_text("\n".join(page.extract_text() or "" for page in reader.pages)).strip()
    if not text:
        raise ValueError("PDF 没有可提取文字，可能是扫描件")
    return text


def normalize_pdf_text(text: str) -> str:
    return text.replace("\x00", "")


def parse_date(text: str) -> date:
    match = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", text)
    if not match:
        match = re.search(r"(?<!\d)(20\d{2})[^\d\r\n]{1,6}(\d{1,2})[^\d\r\n]{1,6}(\d{1,2})(?!\d)", text)
    if not match:
        raise ValueError("未找到开票日期")
    return date(*(int(value) for value in match.groups()))


def parse_invoice_number(text: str) -> str:
    match = re.search(r"(?<!\d)(2\d{19})(?!\d)", text)
    if not match:
        raise ValueError("未找到20位发票号码")
    return match.group(1)


def parse_amount(text: str) -> Decimal:
    compact = re.sub(r"(?<=\d)\s+(?=[\d.,])", "", text)
    compact = re.sub(r"(?<=[.,])\s+(?=\d)", "", compact)
    candidates = re.findall(r"[¥￥´]\s*([0-9][0-9,]*\.\d{2})", compact)
    try:
        amounts = [Decimal(value.replace(",", "")) for value in candidates]
    except InvalidOperation as exc:
        raise ValueError("发票金额格式异常") from exc
    if not amounts:
        raise ValueError("未找到价税合计小写金额")
    return max(amounts)


def is_item_continuation(line: str) -> bool:
    if not line or line.startswith("*"):
        return False
    stop_words = ("合 计", "合        计", "价税合计", "备", "开票人", "发票号码", "开票日期", "订单号", "品牌")
    if line.startswith(stop_words) or "¥" in line or "￥" in line:
        return False
    if re.match(r"^\d+\s*%", line):
        return False
    if len(re.findall(r"\d+(?:\.\d+)?", line)) >= 3:
        return False
    if re.search(r"\s(?:片|个|件|台|条|卷|次|PCS)\s+\d+\s+\d", line, re.IGNORECASE):
        return False
    return True


def clean_item_line(value: str) -> str:
    value = re.split(r"\s+\d+\s*%", value, maxsplit=1)[0]
    value = re.split(
        r"\s+(?:片|个|件|台|条|卷|次|PCS)\s+\d+(?:\.\d+)?\s+\d+(?:\.\d+)?",
        value,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    value = re.split(r"\s+\d+(?:\.\d+)?\s+\d+(?:\.\d+)?", value, maxsplit=1)[0]
    return re.sub(r"\s+", " ", value).strip()


def clean_item_name(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return re.sub(r"\s+无$", "", value)


def parse_item_name(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    items: list[str] = []
    index = 0
    while index < len(lines):
        if not lines[index].startswith("*"):
            index += 1
            continue
        parts = [clean_item_line(lines[index])]
        index += 1
        while index < len(lines) and is_item_continuation(lines[index]):
            parts.append(lines[index])
            index += 1
        item = clean_item_name(" ".join(parts))
        if item and item not in items:
            items.append(item)
    if items and not any(re.search(r"[\x00-\x1f\x7f-\x9f]", item) for item in items):
        return "；".join(items)
    if "通行日期起" in text and re.search(r"(?m)^[ \t]*通行费[ \t]*$", text):
        return "通行费"
    if items:
        raise ValueError("项目名称文字编码异常")
    raise ValueError("未找到项目名称")


def parse_invoice(pdf_path: Path, category: str) -> Invoice:
    text = extract_pdf_text(pdf_path)
    return Invoice(
        category=category,
        source=pdf_path,
        invoice_date=parse_date(text),
        number=parse_invoice_number(text),
        item_name=parse_item_name(text),
        amount=parse_amount(text),
    )


def collect_invoices(root: Path, material_dir: Path, travel_dir: Path) -> tuple[dict[str, list[Invoice]], list[str]]:
    result: dict[str, list[Invoice]] = {"材料费": [], "差旅费": []}
    errors: list[str] = []
    for category, directory in (("材料费", material_dir), ("差旅费", travel_dir)):
        if not directory.is_dir():
            errors.append(f"{directory}: 目录不存在")
            continue
        parsed: list[Invoice] = []
        for pdf_path in sorted(directory.glob("*.pdf")):
            try:
                parsed.append(parse_invoice(pdf_path, category))
            except Exception as exc:  # report every bad file before stopping
                errors.append(f"{pdf_path.relative_to(root)}: {exc}")

        by_number: dict[str, Invoice] = {}
        for invoice in parsed:
            existing = by_number.get(invoice.number)
            if existing is None:
                by_number[invoice.number] = invoice
            elif (existing.invoice_date, existing.item_name, existing.amount) != (
                invoice.invoice_date,
                invoice.item_name,
                invoice.amount,
            ):
                errors.append(f"发票号 {invoice.number} 的重复文件内容不一致")
            else:
                print(f"跳过重复发票：{invoice.source.name}（{invoice.number}）")
        result[category] = sorted(by_number.values(), key=lambda item: (item.invoice_date, item.number))
    return result, errors


def copy_row_format(sheet, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def clear_and_prepare_rows(sheet, count: int, max_column: int) -> None:
    last_row = max(sheet.max_row, 3 + count)
    for row in range(4, last_row + 1):
        if row > sheet.max_row:
            copy_row_format(sheet, 4, row, max_column)
        for column in range(1, max_column + 1):
            sheet.cell(row, column).value = None


def set_text_identifier(cell, value: str) -> None:
    cell.value = value
    cell.number_format = "@"
    cell.quotePrefix = True


def write_material_sheet(sheet, invoices: list[Invoice]) -> None:
    clear_and_prepare_rows(sheet, len(invoices), 10)
    for offset, invoice in enumerate(invoices, start=4):
        sheet.cell(offset, 1, offset - 3)
        sheet.cell(offset, 2, invoice.invoice_date).number_format = "yyyy-mm-dd"
        set_text_identifier(sheet.cell(offset, 3), invoice.number)
        sheet.cell(offset, 4, invoice.item_name)
        sheet.cell(offset, 7, float(invoice.amount)).number_format = "#,##0.00"
        existing_alignment = sheet.cell(offset, 4).alignment
        sheet.cell(offset, 4).alignment = Alignment(
            horizontal=existing_alignment.horizontal or "center",
            vertical="center",
            wrap_text=True,
        )
        if len(invoice.item_name) > 24:
            sheet.row_dimensions[offset].height = max(sheet.row_dimensions[offset].height or 20, 32)


def write_travel_sheet(sheet, invoices: list[Invoice]) -> None:
    clear_and_prepare_rows(sheet, len(invoices), 7)
    sheet.column_dimensions["D"].width = max(sheet.column_dimensions["D"].width or 0, 25)
    for offset, invoice in enumerate(invoices, start=4):
        sheet.cell(offset, 1, offset - 3)
        sheet.cell(offset, 2, invoice.invoice_date).number_format = "yyyy-mm-dd"
        set_text_identifier(sheet.cell(offset, 4), invoice.number)
        sheet.cell(offset, 5, invoice.item_name)
        sheet.cell(offset, 6, float(invoice.amount)).number_format = "#,##0.00"
        sheet.cell(offset, 5).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if len(invoice.item_name) > 12:
            sheet.row_dimensions[offset].height = max(sheet.row_dimensions[offset].height or 20, 30)


def validate_saved_workbook(output_path: Path, invoices: dict[str, list[Invoice]]) -> None:
    workbook = load_workbook(output_path, data_only=False)
    material_numbers = [workbook["材料费"].cell(row, 3).value for row in range(4, 4 + len(invoices["材料费"]))]
    travel_numbers = [workbook["差旅费"].cell(row, 4).value for row in range(4, 4 + len(invoices["差旅费"]))]
    if material_numbers != [invoice.number for invoice in invoices["材料费"]]:
        raise ValueError("保存后的材料费发票号校验失败")
    if travel_numbers != [invoice.number for invoice in invoices["差旅费"]]:
        raise ValueError("保存后的差旅费发票号校验失败")


def write_workbook(template: Path, output_path: Path, invoices: dict[str, list[Invoice]]) -> None:
    workbook = load_workbook(template)
    missing = {"材料费", "差旅费"} - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"台账缺少子表：{', '.join(sorted(missing))}")
    write_material_sheet(workbook["材料费"], invoices["材料费"])
    write_travel_sheet(workbook["差旅费"], invoices["差旅费"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    validate_saved_workbook(output_path, invoices)


def parse_args() -> argparse.Namespace:
    script_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="从材料费、差旅费 PDF 发票生成 IDEA 项目材料台账")
    parser.add_argument("--root", type=Path, default=script_root, help="项目目录")
    parser.add_argument("--workbook", type=Path, help="台账模板路径")
    parser.add_argument("--material-dir", type=Path, help="材料费 PDF 目录")
    parser.add_argument("--travel-dir", type=Path, help="差旅费 PDF 目录")
    parser.add_argument("--output", type=Path, help="输出 xlsx 路径")
    parser.add_argument("--check-only", action="store_true", help="只解析和校验，不生成台账")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = args.root.resolve()
    template = (args.workbook or root / "IDEA项目材料台账.xlsx").resolve()
    material_dir = (args.material_dir or root / "材料费").resolve()
    travel_dir = (args.travel_dir or root / "差旅费").resolve()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = (args.output or root / "输出" / f"IDEA项目材料台账_发票信息统计_{timestamp}.xlsx").resolve()

    if not template.is_file():
        print(f"错误：找不到台账模板 {template}", file=sys.stderr)
        return 1

    invoices, errors = collect_invoices(root, material_dir, travel_dir)
    if errors:
        print("解析未完成，未生成台账：", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1

    material_total = sum((invoice.amount for invoice in invoices["材料费"]), Decimal("0"))
    travel_total = sum((invoice.amount for invoice in invoices["差旅费"]), Decimal("0"))
    print(f"材料费：{len(invoices['材料费'])} 张，合计 {material_total:.2f} 元")
    print(f"差旅费：{len(invoices['差旅费'])} 张，合计 {travel_total:.2f} 元")

    if args.check_only:
        print("校验通过，未生成文件。")
        return 0

    write_workbook(template, output, invoices)
    print(f"已生成：{output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
