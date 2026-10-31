#!/usr/bin/env python3
from __future__ import annotations

import os
import sys
import tempfile
import threading
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import PyQt5

# Qt 在部分中文路径下无法自动定位平台插件。
os.environ.setdefault(
    "QT_QPA_PLATFORM_PLUGIN_PATH",
    str(Path(PyQt5.__file__).resolve().parent / "Qt5" / "plugins" / "platforms"),
)

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment
from pypdf import PdfReader, PdfWriter
from PyQt5.QtCore import QObject, Qt, QThread, pyqtSignal
from PyQt5.QtGui import QColor
from PyQt5.QtPrintSupport import QPrinter, QPrinterInfo
from PyQt5.QtWidgets import (
    QAction,
    QApplication,
    QFileDialog,
    QMainWindow,
    QMessageBox,
    QProgressDialog,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextBrowser,
    QToolBar,
)

from invoice_stats import Invoice, copy_row_format, parse_invoice, set_text_identifier
from pdf_printing import PdfPreviewWidget, PrintSettingsDialog, print_pdf_data
from workbook_printing import workbook_to_pdf_data


FIELD_ALIASES = {
    "sequence": ("序号",),
    "date": ("购买日期", "日期", "开票日期"),
    "number": ("发票号", "发票号码"),
    "item": ("物料品名", "发票名称", "项目名称", "物品名称"),
    "amount": ("金额", "发票金额"),
}

PDF_PREVIEW_TAB = "合并发票预览"
HELP_TAB = "帮助"


def is_invoice_pdf(path: Path) -> bool:
    return path.is_file() and path.suffix.casefold() == ".pdf" and not path.name.startswith("合并发票_")


def discover_invoice_groups(
    invoice_root: Path,
    sheet_names: list[str],
    selected_sheet: str,
) -> dict[str, list[Path]]:
    """按子目录名映射工作表；无含 PDF 子目录时写入当前子表。"""
    if not invoice_root.is_dir():
        raise ValueError(f"发票目录不存在：{invoice_root}")
    if not sheet_names:
        raise ValueError("模板中没有工作表")
    if selected_sheet not in sheet_names:
        raise ValueError("未选中有效的模板子表")

    names = {name.strip().casefold(): name for name in sheet_names}
    child_groups: list[tuple[Path, list[Path]]] = []
    for child in sorted((path for path in invoice_root.iterdir() if path.is_dir()), key=lambda path: path.name):
        pdfs = sorted(
            (path for path in child.rglob("*") if is_invoice_pdf(path)),
            key=lambda path: str(path).casefold(),
        )
        if pdfs:
            child_groups.append((child, pdfs))

    direct_pdfs = sorted(
        (path for path in invoice_root.iterdir() if is_invoice_pdf(path)),
        key=lambda path: str(path).casefold(),
    )
    if not child_groups:
        if not direct_pdfs:
            raise ValueError("所选目录中没有 PDF 发票")
        return {selected_sheet: direct_pdfs}

    if direct_pdfs:
        raise ValueError("发票目录同时含有根目录 PDF 和子目录 PDF，请将根目录 PDF 移入对应子目录")

    groups: dict[str, list[Path]] = {}
    unknown: list[str] = []
    for directory, pdfs in child_groups:
        sheet_name = names.get(directory.name.strip().casefold())
        if sheet_name is None:
            unknown.append(directory.name)
        else:
            groups[sheet_name] = pdfs
    if unknown:
        raise ValueError("以下发票子目录在模板中没有同名子表：" + "、".join(unknown))
    return groups


def find_pdf_files(invoice_root: Path) -> list[Path]:
    if not invoice_root.is_dir():
        raise ValueError(f"发票目录不存在：{invoice_root}")
    pdfs = sorted(
        (path for path in invoice_root.rglob("*") if is_invoice_pdf(path)),
        key=lambda path: str(path.relative_to(invoice_root)).casefold(),
    )
    if not pdfs:
        raise ValueError("所选目录及其子目录中没有 PDF 发票")
    return pdfs


def find_invoice_columns(sheet) -> tuple[int, dict[str, int]]:
    """在前 30 行中定位包含日期、发票号、名称和金额的表头。"""
    required = {"date", "number", "item", "amount"}
    for row in range(1, min(sheet.max_row, 30) + 1):
        columns: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row, column).value
            if value is None:
                continue
            label = str(value).strip()
            for field, aliases in FIELD_ALIASES.items():
                if label in aliases and field not in columns:
                    columns[field] = column
        if required <= columns.keys():
            return row, columns
    raise ValueError(f"子表“{sheet.title}”未找到完整表头（日期、发票号、名称、金额）")


def deduplicate_invoices(invoices: list[Invoice]) -> tuple[list[Invoice], int]:
    by_number: dict[str, Invoice] = {}
    duplicate_count = 0
    for invoice in invoices:
        existing = by_number.get(invoice.number)
        if existing is None:
            by_number[invoice.number] = invoice
        elif (existing.invoice_date, existing.item_name, existing.amount) == (
            invoice.invoice_date,
            invoice.item_name,
            invoice.amount,
        ):
            duplicate_count += 1
        else:
            raise ValueError(f"发票号 {invoice.number} 的重复文件内容不一致")
    return sorted(by_number.values(), key=lambda item: (item.invoice_date, item.number)), duplicate_count


def format_statistics_message(result: dict[str, list[Invoice]], duplicates: int) -> str:
    count = sum(len(invoices) for invoices in result.values())
    total = sum((invoice.amount for invoices in result.values() for invoice in invoices), Decimal("0.00"))
    suffix = f"，跳过{duplicates}张重复发票" if duplicates else ""
    return f"统计完成：共{count}张发票，总金额为：{total:.2f}{suffix}"


def write_invoices_to_sheet(sheet, invoices: list[Invoice], header_row: int, columns: dict[str, int]) -> None:
    first_data_row = header_row + 1
    template_row = first_data_row
    previous_last_row = sheet.max_row
    target_last_row = first_data_row + len(invoices) - 1
    clear_last_row = max(previous_last_row, target_last_row)

    for row in range(first_data_row, clear_last_row + 1):
        if row > previous_last_row:
            copy_row_format(sheet, template_row, row, sheet.max_column)
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)
            if not isinstance(cell, MergedCell):
                cell.value = None

    for row, invoice in enumerate(invoices, start=first_data_row):
        if "sequence" in columns:
            sheet.cell(row, columns["sequence"], row - header_row)
        sheet.cell(row, columns["date"], invoice.invoice_date).number_format = "yyyy-mm-dd"
        set_text_identifier(sheet.cell(row, columns["number"]), invoice.number)
        item_cell = sheet.cell(row, columns["item"], invoice.item_name)
        amount_cell = sheet.cell(row, columns["amount"], float(invoice.amount))
        amount_cell.number_format = "#,##0.00"
        item_cell.alignment = Alignment(
            horizontal=item_cell.alignment.horizontal or "center",
            vertical="center",
            wrap_text=True,
        )
        if len(invoice.item_name) > 24:
            sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 20, 32)


def apply_invoice_results(workbook, result: dict[str, list[Invoice]]):
    snapshot = BytesIO()
    workbook.save(snapshot)
    snapshot.seek(0)
    updated = load_workbook(snapshot, data_only=False)
    layouts = {name: find_invoice_columns(updated[name]) for name in result}
    for sheet_name, invoices in result.items():
        header_row, columns = layouts[sheet_name]
        write_invoices_to_sheet(updated[sheet_name], invoices, header_row, columns)
    return updated


class ExtractionWorker(QObject):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(object, object, int)
    cancelled = pyqtSignal()

    def __init__(self, groups: dict[str, list[Path]], cancel_event: threading.Event):
        super().__init__()
        self.groups = groups
        self.cancel_event = cancel_event

    def run(self) -> None:
        total = sum(len(paths) for paths in self.groups.values())
        completed = 0
        errors: list[str] = []
        result: dict[str, list[Invoice]] = {}
        duplicate_count = 0
        for sheet_name, paths in self.groups.items():
            parsed: list[Invoice] = []
            for pdf_path in paths:
                if self.cancel_event.is_set():
                    self.cancelled.emit()
                    return
                try:
                    parsed.append(parse_invoice(pdf_path, sheet_name))
                except Exception as exc:
                    errors.append(f"{pdf_path.name}：{exc}")
                completed += 1
                self.progress.emit(round(completed * 100 / total), pdf_path.name)
            try:
                result[sheet_name], skipped = deduplicate_invoices(parsed)
                duplicate_count += skipped
            except ValueError as exc:
                errors.append(f"{sheet_name}：{exc}")
        self.finished.emit(result, errors, duplicate_count)


class MergeWorker(QObject):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(object, int, int)
    failed = pyqtSignal(str)
    cancelled = pyqtSignal()

    def __init__(self, pdf_paths: list[Path], cancel_event: threading.Event):
        super().__init__()
        self.pdf_paths = pdf_paths
        self.cancel_event = cancel_event

    def run(self) -> None:
        writer = PdfWriter()
        readers: list[PdfReader] = []
        temp_path: Path | None = None
        expected_pages = 0
        try:
            for index, pdf_path in enumerate(self.pdf_paths, start=1):
                if self.cancel_event.is_set():
                    self.cancelled.emit()
                    return
                reader = PdfReader(str(pdf_path))
                readers.append(reader)
                for page in reader.pages:
                    if self.cancel_event.is_set():
                        self.cancelled.emit()
                        return
                    writer.add_page(page)
                    expected_pages += 1
                self.progress.emit(round(index * 90 / len(self.pdf_paths)), pdf_path.name)

            if self.cancel_event.is_set():
                self.cancelled.emit()
                return
            with tempfile.NamedTemporaryFile(prefix="merged-invoices-", suffix=".pdf", delete=False) as stream:
                temp_path = Path(stream.name)
            self.progress.emit(95, "正在写入合并文件")
            with temp_path.open("wb") as stream:
                writer.write(stream)
            if self.cancel_event.is_set():
                temp_path.unlink(missing_ok=True)
                self.cancelled.emit()
                return

            checked = PdfReader(str(temp_path))
            actual_pages = len(checked.pages)
            checked.close()
            if actual_pages != expected_pages:
                raise ValueError(f"合并页数校验失败：应为 {expected_pages} 页，实际为 {actual_pages} 页")
            self.progress.emit(100, "合并完成")
            self.finished.emit(temp_path, len(self.pdf_paths), actual_pages)
        except Exception as exc:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
            self.failed.emit(str(exc))
        finally:
            for reader in readers:
                reader.close()
            writer.close()


class InvoiceWindow(QMainWindow):
    def __init__(self, root: Path | None = None):
        super().__init__()
        default_root = Path.home() if getattr(sys, "frozen", False) else Path(__file__).resolve().parents[1]
        self.root = (root or default_root).resolve()
        self.template_path: Path | None = None
        self.invoice_root: Path | None = None
        self.workbook = None
        self.tables: dict[QTableWidget, str] = {}
        self.loading_tables = False
        self.thread: QThread | None = None
        self.worker: QObject | None = None
        self.cancel_event: threading.Event | None = None
        self.progress_dialog: QProgressDialog | None = None
        self.merged_pdf_data: bytes | None = None
        self.closed_sheet_names: set[str] = set()
        self.pdf_preview_open = False
        self.help_open = False
        self.busy = False

        self.setWindowTitle("发票整理工具 v1.1.2")
        self.resize(1180, 720)
        self.tabs = QTabWidget()
        self.tabs.setTabsClosable(True)
        self.setCentralWidget(self.tabs)
        self._build_toolbar()
        self.tabs.tabCloseRequested.connect(self._close_workspace_tab)
        self.tabs.currentChanged.connect(lambda _index: self._update_actions())
        self.statusBar().showMessage("请先选择 Excel 模板和发票目录")
        self._update_actions()

    def _build_toolbar(self) -> None:
        toolbar = QToolBar("工具栏")
        toolbar.setMovable(False)
        self.addToolBar(Qt.TopToolBarArea, toolbar)
        self.template_action = QAction("选择模板", self)
        self.invoice_action = QAction("选择发票", self)
        self.stats_action = QAction("统计", self)
        self.merge_action = QAction("合并", self)
        self.export_action = QAction("导出", self)
        self.print_action = QAction("打印", self)
        self.help_action = QAction("帮助", self)
        self.exit_action = QAction("退出", self)
        for action in (
            self.template_action,
            self.invoice_action,
            self.stats_action,
            self.merge_action,
            self.export_action,
            self.print_action,
            self.help_action,
            self.exit_action,
        ):
            toolbar.addAction(action)
        self.template_action.triggered.connect(self.choose_template)
        self.invoice_action.triggered.connect(self.choose_invoices)
        self.stats_action.triggered.connect(self.start_statistics)
        self.merge_action.triggered.connect(self.start_merge)
        self.export_action.triggered.connect(self.export_current)
        self.print_action.triggered.connect(self.start_print)
        self.help_action.triggered.connect(self.show_help)
        self.exit_action.triggered.connect(self.close)

    def _update_actions(self, busy: bool | None = None) -> None:
        if busy is not None:
            self.busy = busy
        workspace_kind = self._current_workspace_kind()
        self.template_action.setEnabled(not self.busy)
        self.invoice_action.setEnabled(not self.busy)
        self.stats_action.setEnabled(
            not self.busy
            and workspace_kind == "workbook"
            and self.workbook is not None
            and self.invoice_root is not None
        )
        self.merge_action.setEnabled(not self.busy and self.invoice_root is not None)
        self.export_action.setEnabled(not self.busy and workspace_kind in {"workbook", "pdf"})
        self.print_action.setEnabled(not self.busy and workspace_kind in {"workbook", "pdf"})
        self.help_action.setEnabled(not self.busy)
        self.exit_action.setEnabled(not self.busy)
        self.tabs.setEnabled(not self.busy)

    def _current_workspace_kind(self) -> str | None:
        widget = self.tabs.currentWidget()
        return widget.property("workspaceKind") if widget is not None else None

    def choose_template(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(
            self,
            "选择 Excel 模板",
            str(self.root),
            "Excel 工作簿 (*.xlsx)",
        )
        if not filename:
            return
        try:
            workbook = load_workbook(filename, data_only=False)
        except Exception as exc:
            self._show_error(f"模板打开失败：{exc}")
            return
        self.template_path = Path(filename)
        self.root = self.template_path.parent
        self.workbook = workbook
        self.closed_sheet_names.clear()
        self._render_workbook()
        self._update_actions()
        self.statusBar().showMessage(f"已导入模板：{self.template_path.name}")

    def choose_invoices(self) -> None:
        directory = QFileDialog.getExistingDirectory(self, "选择发票目录", str(self.root))
        if not directory:
            return
        self.invoice_root = Path(directory)
        self._update_actions()
        self.statusBar().showMessage(f"已选择发票目录：{self.invoice_root}")

    def start_statistics(self) -> None:
        if self.workbook is None or self.template_path is None or self.invoice_root is None:
            return
        try:
            selected_sheet = self.tables.get(self.tabs.currentWidget())
            if selected_sheet is None:
                raise ValueError("请先选择一个 Excel 子表")
            groups = discover_invoice_groups(self.invoice_root, self.workbook.sheetnames, selected_sheet)
            for sheet_name in groups:
                find_invoice_columns(self.workbook[sheet_name])
        except Exception as exc:
            self._show_error(str(exc))
            return

        total = sum(len(paths) for paths in groups.values())
        self.cancel_event = threading.Event()
        self.progress_dialog = QProgressDialog("正在读取发票…", "取消", 0, 100, self)
        self.progress_dialog.setWindowTitle(f"统计发票（共 {total} 张）")
        self.progress_dialog.setWindowModality(Qt.WindowModal)
        self.progress_dialog.setAutoClose(False)
        self.progress_dialog.setAutoReset(False)
        self.progress_dialog.canceled.connect(self.cancel_event.set)
        self.progress_dialog.show()
        self._update_actions(busy=True)

        self.thread = QThread(self)
        self.worker = ExtractionWorker(groups, self.cancel_event)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.progress.connect(self._on_progress)
        self.worker.finished.connect(self._on_statistics_finished)
        self.worker.cancelled.connect(self._on_cancelled)
        self.worker.finished.connect(self.thread.quit)
        self.worker.cancelled.connect(self.thread.quit)
        self.thread.finished.connect(self._clear_worker)
        self.thread.start()

    def _on_progress(self, value: int, filename: str) -> None:
        dialog = self.progress_dialog
        if dialog is None:
            return
        dialog.setLabelText(f"正在读取：{filename}")
        self.statusBar().showMessage(f"统计进度 {value}%：{filename}")
        # setValue 可能处理完成事件并关闭进度窗口，因此必须放在末尾。
        dialog.setValue(value)

    def _on_statistics_finished(self, result: dict[str, list[Invoice]], errors: list[str], duplicates: int) -> None:
        self._close_progress()
        if errors:
            preview = "\n".join(f"- {error}" for error in errors[:15])
            if len(errors) > 15:
                preview += f"\n- 其余 {len(errors) - 15} 项…"
            self._show_error("发票解析未完成，工作区未修改：\n" + preview)
            return
        try:
            updated = apply_invoice_results(self.workbook, result)
        except Exception as exc:
            self._show_error(f"写入工作区失败：{exc}")
            return

        self.workbook = updated
        self._render_workbook()
        self.statusBar().showMessage(format_statistics_message(result, duplicates))

    def _on_cancelled(self) -> None:
        self._close_progress()
        self.statusBar().showMessage("已取消统计，工作区未修改")

    def start_merge(self) -> None:
        if self.invoice_root is None:
            return
        try:
            pdf_paths = find_pdf_files(self.invoice_root)
        except Exception as exc:
            self._show_error(str(exc))
            return

        self.cancel_event = threading.Event()
        self.progress_dialog = QProgressDialog("正在合并发票…", "取消", 0, 100, self)
        self.progress_dialog.setWindowTitle(f"合并发票（共 {len(pdf_paths)} 个 PDF）")
        self.progress_dialog.setWindowModality(Qt.WindowModal)
        self.progress_dialog.setAutoClose(False)
        self.progress_dialog.setAutoReset(False)
        self.progress_dialog.canceled.connect(self.cancel_event.set)
        self.progress_dialog.show()
        self._update_actions(busy=True)

        self.thread = QThread(self)
        self.worker = MergeWorker(pdf_paths, self.cancel_event)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.progress.connect(self._on_merge_progress)
        self.worker.finished.connect(self._on_merge_finished)
        self.worker.failed.connect(self._on_merge_failed)
        self.worker.cancelled.connect(self._on_merge_cancelled)
        self.worker.finished.connect(self.thread.quit)
        self.worker.failed.connect(self.thread.quit)
        self.worker.cancelled.connect(self.thread.quit)
        self.thread.finished.connect(self._clear_worker)
        self.thread.start()

    def _on_merge_progress(self, value: int, filename: str) -> None:
        dialog = self.progress_dialog
        if dialog is None:
            return
        dialog.setLabelText(f"正在合并：{filename}")
        self.statusBar().showMessage(f"合并进度 {value}%：{filename}")
        dialog.setValue(value)

    def _on_merge_finished(self, temp_path: Path, file_count: int, page_count: int) -> None:
        self._close_progress()
        try:
            self.merged_pdf_data = temp_path.read_bytes()
            self._show_pdf_preview()
        except Exception as exc:
            temp_path.unlink(missing_ok=True)
            self.merged_pdf_data = None
            self._update_actions()
            self._show_error(f"显示合并发票失败：{exc}")
            return
        finally:
            temp_path.unlink(missing_ok=True)
        self.statusBar().showMessage(
            f"合并完成：{file_count} 个 PDF，共 {page_count} 页；请点击“导出”保存合并文件"
        )

    def _show_pdf_preview(self) -> None:
        if self.merged_pdf_data is None:
            return
        for index in range(self.tabs.count()):
            if self.tabs.tabText(index) == PDF_PREVIEW_TAB:
                old_widget = self.tabs.widget(index)
                self.tabs.removeTab(index)
                old_widget.deleteLater()
                break
        preview = PdfPreviewWidget(self.merged_pdf_data)
        preview.setProperty("workspaceKind", "pdf")
        self.tabs.addTab(preview, PDF_PREVIEW_TAB)
        self.tabs.setCurrentWidget(preview)
        self.pdf_preview_open = True
        self._update_actions()

    def start_print(self) -> None:
        workspace_kind = self._current_workspace_kind()
        if workspace_kind == "pdf":
            pdf_data = self.merged_pdf_data
            document_name = "合并发票"
        elif workspace_kind == "workbook" and self.workbook is not None:
            QApplication.setOverrideCursor(Qt.WaitCursor)
            self.statusBar().showMessage("正在生成 Excel 打印预览…")
            try:
                pdf_data = workbook_to_pdf_data(self.workbook)
            except Exception as exc:
                self._show_error(f"生成 Excel 打印预览失败：{exc}")
                return
            finally:
                QApplication.restoreOverrideCursor()
            document_name = self.template_path.name if self.template_path else "Excel 工作簿"
        else:
            return
        if pdf_data is None:
            return
        dialog = PrintSettingsDialog(pdf_data, self)
        if dialog.exec_() != dialog.Accepted:
            return
        settings = dialog.settings()
        printer_info = QPrinterInfo.printerInfo(settings.printer_name)
        printer = QPrinter(printer_info, QPrinter.HighResolution)
        resolutions = [value for value in printer_info.supportedResolutions() if value <= 300]
        if resolutions:
            printer.setResolution(max(resolutions))
        self._update_actions(busy=True)
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            sheet_count = print_pdf_data(pdf_data, settings, printer)
        except Exception as exc:
            self._show_error(f"打印失败：{exc}")
            return
        finally:
            QApplication.restoreOverrideCursor()
            self._update_actions()
        self.statusBar().showMessage(f"已将{document_name}提交到 {settings.printer_name}，共 {sheet_count} 张纸")
        QMessageBox.information(self, "已提交打印", f"打印任务已发送到：\n{settings.printer_name}")

    def _on_merge_failed(self, message: str) -> None:
        self._close_progress()
        self._show_error(f"合并失败：{message}")

    def _on_merge_cancelled(self) -> None:
        self._close_progress()
        self.statusBar().showMessage("已取消合并，未生成文件")

    def _close_progress(self) -> None:
        if self.progress_dialog is not None:
            self.progress_dialog.close()
            self.progress_dialog = None
        self._update_actions(busy=False)

    def _clear_worker(self) -> None:
        if self.worker is not None:
            self.worker.deleteLater()
        if self.thread is not None:
            self.thread.deleteLater()
        self.worker = None
        self.thread = None
        self.cancel_event = None

    def _render_workbook(self) -> None:
        selected_sheet = self.tabs.tabText(self.tabs.currentIndex()) if self.tabs.count() else ""
        self.loading_tables = True
        while self.tabs.count():
            widget = self.tabs.widget(0)
            self.tabs.removeTab(0)
            widget.deleteLater()
        self.tables.clear()
        for sheet in self.workbook.worksheets:
            if sheet.title in self.closed_sheet_names:
                continue
            table = QTableWidget(max(sheet.max_row, 1), max(sheet.max_column, 1))
            table.setProperty("workspaceKind", "workbook")
            table.setAlternatingRowColors(True)
            table.setHorizontalHeaderLabels([self._column_name(index) for index in range(1, table.columnCount() + 1)])
            table.setVerticalHeaderLabels([str(index) for index in range(1, table.rowCount() + 1)])
            table.itemChanged.connect(self._table_item_changed)
            merged_children: set[tuple[int, int]] = set()
            for merged_range in sheet.merged_cells.ranges:
                min_col, min_row, max_col, max_row = merged_range.bounds
                table.setSpan(min_row - 1, min_col - 1, max_row - min_row + 1, max_col - min_col + 1)
                for row in range(min_row, max_row + 1):
                    for column in range(min_col, max_col + 1):
                        if (row, column) != (min_row, min_col):
                            merged_children.add((row, column))
            for row in range(1, table.rowCount() + 1):
                height = sheet.row_dimensions[row].height
                if height:
                    table.setRowHeight(row - 1, max(20, int(height * 1.35)))
                for column in range(1, table.columnCount() + 1):
                    cell = sheet.cell(row, column)
                    value = "" if cell.value is None else self._display_value(cell.value)
                    item = QTableWidgetItem(value)
                    item.setData(Qt.UserRole, self._value_kind(cell.value))
                    if (row, column) in merged_children:
                        item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                    if cell.fill.fill_type and cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb:
                        color = cell.fill.fgColor.rgb[-6:]
                        item.setBackground(QColor(f"#{color}"))
                    table.setItem(row - 1, column - 1, item)
            for column in range(1, table.columnCount() + 1):
                width = sheet.column_dimensions[self._column_name(column)].width
                longest = max(
                    (len(table.item(row, column - 1).text()) for row in range(table.rowCount())),
                    default=0,
                )
                table.setColumnWidth(
                    column - 1,
                    min(380, max(80, int((width or 12) * 8), longest * 10 + 32)),
                )
            self.tabs.addTab(table, sheet.title)
            self.tables[table] = sheet.title
        for index in range(self.tabs.count()):
            if self.tabs.tabText(index) == selected_sheet:
                self.tabs.setCurrentIndex(index)
                break
        self.loading_tables = False
        if self.pdf_preview_open and self.merged_pdf_data is not None:
            preview = PdfPreviewWidget(self.merged_pdf_data)
            preview.setProperty("workspaceKind", "pdf")
            self.tabs.addTab(preview, PDF_PREVIEW_TAB)
            if selected_sheet == PDF_PREVIEW_TAB:
                self.tabs.setCurrentWidget(preview)
        if self.help_open:
            help_widget = self._create_help_widget()
            self.tabs.addTab(help_widget, HELP_TAB)
            if selected_sheet == HELP_TAB:
                self.tabs.setCurrentWidget(help_widget)
        self._update_actions()

    def _table_item_changed(self, item: QTableWidgetItem) -> None:
        if self.loading_tables or self.workbook is None:
            return
        table = item.tableWidget()
        sheet_name = self.tables.get(table)
        if sheet_name is None:
            return
        cell = self.workbook[sheet_name].cell(item.row() + 1, item.column() + 1)
        try:
            cell.value = self._parse_edited_value(item.text(), item.data(Qt.UserRole))
            item.setData(Qt.UserRole, self._value_kind(cell.value))
            self.statusBar().showMessage(f"已修改 {sheet_name}!{cell.coordinate}")
        except ValueError as exc:
            self._show_error(str(exc))
            self.loading_tables = True
            item.setText(self._display_value(cell.value))
            self.loading_tables = False

    def export_current(self) -> None:
        workspace_kind = self._current_workspace_kind()
        if workspace_kind == "pdf":
            self.export_pdf()
        elif workspace_kind == "workbook":
            self.export_workbook()

    def export_pdf(self) -> None:
        if self.merged_pdf_data is None:
            return
        output_dir = self.root / "输出"
        output_dir.mkdir(parents=True, exist_ok=True)
        default_name = f"合并发票_{datetime.now():%Y%m%d_%H%M}.pdf"
        filename, _ = QFileDialog.getSaveFileName(
            self,
            "导出合并发票",
            str(output_dir / default_name),
            "PDF 文件 (*.pdf)",
        )
        if not filename:
            return
        output_path = Path(filename)
        if output_path.suffix.lower() != ".pdf":
            output_path = output_path.with_suffix(".pdf")
        temp_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix="invoice-export-",
                suffix=".pdf",
                dir=output_path.parent,
                delete=False,
            ) as stream:
                stream.write(self.merged_pdf_data)
                temp_path = Path(stream.name)
            checked = PdfReader(str(temp_path))
            page_count = len(checked.pages)
            checked.close()
            if page_count == 0:
                raise ValueError("合并发票没有页面")
            os.replace(temp_path, output_path)
        except Exception as exc:
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)
            self._show_error(f"导出 PDF 失败：{exc}")
            return
        self.statusBar().showMessage(f"已导出：{output_path}")
        QMessageBox.information(self, "导出完成", f"已保存到：\n{output_path}")

    def export_workbook(self) -> None:
        if self.workbook is None or self.template_path is None:
            return
        output_dir = self.root / "输出"
        output_dir.mkdir(parents=True, exist_ok=True)
        default_name = f"{self.template_path.stem}_{datetime.now():%Y%m%d_%H%M}.xlsx"
        filename, _ = QFileDialog.getSaveFileName(
            self,
            "导出 Excel",
            str(output_dir / default_name),
            "Excel 工作簿 (*.xlsx)",
        )
        if not filename:
            return
        output_path = Path(filename)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        try:
            self.workbook.save(output_path)
            checked = load_workbook(output_path, read_only=True, data_only=False)
            if checked.sheetnames != self.workbook.sheetnames:
                raise ValueError("导出文件的子表校验失败")
            checked.close()
        except Exception as exc:
            self._show_error(f"导出失败：{exc}")
            return
        self.statusBar().showMessage(f"已导出：{output_path}")
        QMessageBox.information(self, "导出完成", f"已保存到：\n{output_path}")

    def _close_workspace_tab(self, index: int) -> None:
        widget = self.tabs.widget(index)
        if widget is None:
            return
        workspace_kind = widget.property("workspaceKind")
        if workspace_kind == "workbook":
            sheet_name = self.tables.pop(widget, None)
            if sheet_name:
                self.closed_sheet_names.add(sheet_name)
        elif workspace_kind == "pdf":
            self.pdf_preview_open = False
            self.merged_pdf_data = None
        elif workspace_kind == "help":
            self.help_open = False
        self.tabs.removeTab(index)
        widget.deleteLater()
        self._update_actions()

    def show_help(self) -> None:
        for index in range(self.tabs.count()):
            if self.tabs.tabText(index) == HELP_TAB:
                self.tabs.setCurrentIndex(index)
                return
        help_widget = self._create_help_widget()
        self.tabs.addTab(help_widget, HELP_TAB)
        self.tabs.setCurrentWidget(help_widget)
        self.help_open = True
        self._update_actions()

    @staticmethod
    def _create_help_widget() -> QTextBrowser:
        help_widget = QTextBrowser()
        help_widget.setProperty("workspaceKind", "help")
        help_widget.setOpenExternalLinks(True)
        help_widget.setHtml(
            """
            <style>
              body { font-family: 'PingFang SC', 'Microsoft YaHei', sans-serif; margin: 28px; line-height: 1.65; }
              h1 { color: #1f4e78; }
              h2 { margin-top: 24px; color: #333; }
              code { background: #f2f2f2; padding: 2px 5px; }
            </style>
            <h1>发票整理工具 v1.1.2</h1>
            <p>本工具用于从电子发票 PDF 中提取信息、写入 Excel 台账、合并发票、预览、导出和打印。</p>

            <h2>基本操作</h2>
            <ol>
              <li>点击“选择模板”，导入 <code>IDEA项目材料台账.xlsx</code>。</li>
              <li>点击“选择发票”，选择包含发票 PDF 的目录。</li>
              <li>在工作区选择目标 Excel 子表，点击“统计”提取发票日期、号码、项目名称和含税金额。</li>
              <li>点击“合并”，将所选目录及子目录中的发票合并，并在“合并发票预览”标签中查看。</li>
            </ol>

            <h2>导出与打印</h2>
            <ul>
              <li>当前显示 Excel 子表时，“导出”保存整个 Excel 工作簿，“打印”打印工作簿。</li>
              <li>当前显示“合并发票预览”时，“导出”保存合并 PDF，“打印”打印合并发票。</li>
              <li>打印窗口支持页面大小和一张多页两种方式。确认后任务会直接发送到所选系统打印机。</li>
              <li>标签右侧的“×”可关闭对应工作区内容。关闭合并发票预览后，需要重新合并才能再次导出或打印。</li>
            </ul>

            <h2>版本与授权</h2>
            <p><b>版本：</b>v1.1.2。</p>
            <p><b>非商业使用：</b>个人学习、研究和其他非商业用途可免费使用。</p>
            <p><b>商业使用：</b>用于企业经营、收费服务、对外交付或其他商业用途时，需要付费获得商业授权。具体授权请联系软件提供方。</p>
            """
        )
        return help_widget

    def closeEvent(self, event) -> None:
        if self.thread is not None and self.thread.isRunning():
            event.ignore()
            return
        event.accept()

    def _show_error(self, message: str) -> None:
        self.statusBar().showMessage(message.replace("\n", " "))
        QMessageBox.critical(self, "出错了", message)

    @staticmethod
    def _column_name(index: int) -> str:
        name = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            name = chr(65 + remainder) + name
        return name

    @staticmethod
    def _display_value(value) -> str:
        if value is None:
            return ""
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, float):
            return f"{value:g}"
        return str(value)

    @staticmethod
    def _value_kind(value) -> str:
        if isinstance(value, (date, datetime)):
            return "date"
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            return "number"
        return "text"

    @staticmethod
    def _parse_edited_value(text: str, kind: str):
        text = text.strip()
        if not text:
            return None
        if text.startswith("="):
            return text
        if kind == "date":
            for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
                try:
                    return datetime.strptime(text, pattern).date()
                except ValueError:
                    pass
            raise ValueError("日期格式应为 YYYY-MM-DD")
        if kind == "number":
            try:
                return float(text.replace(",", ""))
            except ValueError as exc:
                raise ValueError("数值格式不正确") from exc
        return text


def main() -> int:
    app = QApplication(sys.argv)
    app.setApplicationName("发票整理工具")
    window = InvoiceWindow()
    window.show()
    return app.exec_()


if __name__ == "__main__":
    raise SystemExit(main())
