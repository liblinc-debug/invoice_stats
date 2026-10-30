import tempfile
import threading
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from gui import (
    ExtractionWorker,
    InvoiceWindow,
    MergeWorker,
    apply_invoice_results,
    discover_invoice_groups,
    find_pdf_files,
    find_invoice_columns,
    format_statistics_message,
    write_invoices_to_sheet,
)
from invoice_stats import Invoice


class GuiLogicTest(unittest.TestCase):
    def test_merge_worker_combines_recursive_pdfs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            nested = root / "差旅费"
            nested.mkdir()
            for path, pages in (
                (root / "a.pdf", 2),
                (nested / "b.PDF", 1),
                (root / "合并发票_20261030_1642.pdf", 1),
            ):
                writer = PdfWriter()
                for _ in range(pages):
                    writer.add_blank_page(width=72, height=72)
                with path.open("wb") as stream:
                    writer.write(stream)
                writer.close()

            pdf_paths = find_pdf_files(root)
            self.assertEqual(pdf_paths, [root / "a.pdf", nested / "b.PDF"])
            completed = []
            errors = []
            worker = MergeWorker(pdf_paths, threading.Event())
            worker.finished.connect(lambda *values: completed.append(values))
            worker.failed.connect(errors.append)
            worker.run()

            self.assertEqual(errors, [])
            self.assertEqual(len(completed), 1)
            merged_path, file_count, page_count = completed[0]
            self.assertEqual((file_count, page_count), (2, 3))
            merged = PdfReader(str(merged_path))
            self.assertEqual(len(merged.pages), 3)
            merged.close()
            Path(merged_path).unlink()

    def test_merge_worker_can_cancel_before_reading(self):
        event = threading.Event()
        event.set()
        worker = MergeWorker([Path("not-read.pdf")], event)
        cancelled = []
        worker.cancelled.connect(lambda: cancelled.append(True))
        worker.run()
        self.assertEqual(cancelled, [True])

    def test_single_sheet_statistics_preserves_other_sheets(self):
        workbook = Workbook()
        material = workbook.active
        material.title = "材料费"
        material["A1"] = "保留的材料费数据"
        material["B2"] = "=1+1"
        travel = workbook.create_sheet("差旅费")
        travel.append(["序号", "日期", "发票号", "发票名称", "发票金额"])
        invoice = Invoice("差旅费", Path("trip.pdf"), date(2026, 2, 3), "123", "住宿费", Decimal("88.50"))

        updated = apply_invoice_results(workbook, {"差旅费": [invoice]})

        self.assertEqual(updated["材料费"]["A1"].value, "保留的材料费数据")
        self.assertEqual(updated["材料费"]["B2"].value, "=1+1")
        self.assertEqual(updated["差旅费"]["C2"].value, "123")
        self.assertEqual(updated["差旅费"]["E2"].value, 88.5)

    def test_statistics_message_includes_decimal_total(self):
        invoices = {
            "材料费": [Invoice("材料费", Path("a.pdf"), date(2026, 1, 1), "1", "A", Decimal("45.70"))],
            "差旅费": [Invoice("差旅费", Path("b.pdf"), date(2026, 1, 2), "2", "B", Decimal("0.30"))],
        }
        self.assertEqual(
            format_statistics_message(invoices, 1),
            "统计完成：共2张发票，总金额为：46.00，跳过1张重复发票",
        )

    def test_progress_completion_can_close_dialog_during_set_value(self):
        class StatusBar:
            message = ""

            def showMessage(self, message):
                self.message = message

        class Window:
            def __init__(self):
                self.bar = StatusBar()
                self.progress_dialog = Dialog(self)

            def statusBar(self):
                return self.bar

        class Dialog:
            def __init__(self, window):
                self.window = window
                self.label = ""

            def setLabelText(self, label):
                self.label = label

            def setValue(self, _value):
                self.window.progress_dialog = None

        window = Window()
        InvoiceWindow._on_progress(window, 100, "last.pdf")
        self.assertIsNone(window.progress_dialog)
        self.assertEqual(window.bar.message, "统计进度 100%：last.pdf")

    def test_worker_can_cancel_before_parsing(self):
        event = threading.Event()
        event.set()
        worker = ExtractionWorker({"材料费": [Path("not-read.pdf")]}, event)
        cancelled = []
        worker.cancelled.connect(lambda: cancelled.append(True))
        worker.run()
        self.assertEqual(cancelled, [True])

    def test_directory_mapping_and_selected_sheet_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "one.pdf").touch()
            self.assertEqual(
                discover_invoice_groups(root, ["材料费", "差旅费"], "差旅费"),
                {"差旅费": [root / "one.pdf"]},
            )

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            child = root / "差旅费"
            child.mkdir()
            (child / "trip.pdf").touch()
            self.assertEqual(
                discover_invoice_groups(root, ["材料费", "差旅费"], "材料费"),
                {"差旅费": [child / "trip.pdf"]},
            )

    def test_header_driven_write_and_sortable_types(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "材料费"
        sheet.append(["标题", None, None, None, None])
        sheet.append([None, None, None, None, None])
        sheet.append(["序号", "发票号", "金额", "购买日期", "物品名称"])
        sheet.append([99, "old", 1, date(2020, 1, 1), "old"])
        invoice = Invoice("材料费", Path("a.pdf"), date(2026, 2, 3), "26952000000168054676", "摄像头", Decimal("45.70"))

        header_row, columns = find_invoice_columns(sheet)
        write_invoices_to_sheet(sheet, [invoice], header_row, columns)

        self.assertEqual(sheet["A4"].value, 1)
        self.assertEqual(sheet["B4"].value, invoice.number)
        self.assertEqual(sheet["C4"].value, 45.7)
        self.assertEqual(sheet["D4"].value, invoice.invoice_date)
        self.assertEqual(sheet["E4"].value, invoice.item_name)


if __name__ == "__main__":
    unittest.main()
